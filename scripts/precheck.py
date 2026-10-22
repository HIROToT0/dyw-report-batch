# -*- coding: utf-8 -*-
"""批量预检: 不写入任何数据, 只判断 85 份能否自动跑完。"""
import sys, os, json
from collections import Counter, defaultdict
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dyw_client import DywClient, parse_volume, parse_temperature, format_remark_a
from report_engine import SIZE_RANGE, LOAD_RANGE

XLSX = os.getenv("DYW_LEDGER", r'C:/Users/bonbo/Desktop/未发报告.xlsx')


def col_of(ws, *names, default=None):
    """按表头名定位列(1-based)，兼容各单位导出模板列序不同。"""
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or "").strip()
        if h in names:
            return c
    return default


def load():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['任务列表']
    c_temp = col_of(ws, "温度", default=6)
    c_hum = col_of(ws, "湿度", default=7)
    print("台账列定位: 温度=第%d列 湿度=第%d列" % (c_temp, c_hum))
    g = defaultdict(lambda: {"samples": [], "temp": None, "hum": None})
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if not code: continue
        t = g[str(code).strip()]
        s = ws.cell(r, 2).value
        if s: t["samples"].append(str(s).strip())
        t["temp"] = ws.cell(r, c_temp).value
        t["hum"] = ws.cell(r, c_hum).value
    return g

c = DywClient(); c.login()
tasks = load()
ok, warn = [], []
for i, (code, t) in enumerate(sorted(tasks.items()), 1):
    j = c.post('/ptcompile/getReportInfo.do', data={'reportCode': code, 'reportId': '551'}).json()
    if not j.get('isSuccess'):
        warn.append((code, "查询失败", str(j.get('responseMsg')))); continue
    d = j['data']; ss = d['savedSingleField']
    gui, grade = ss.get('gui_ge'), ss.get('chanpin_guige')
    vol = parse_volume(ss.get('qi_ta'))
    probs = []
    if gui not in SIZE_RANGE: probs.append("规格不支持:%s" % gui)
    if grade not in LOAD_RANGE: probs.append("等级不支持:%s" % grade)
    if vol is None: probs.append("方量解析失败:%r" % (ss.get('qi_ta') or ''))
    yh = (ss.get('yang_hu_tiao_jian') or '').strip()
    temp = None
    if yh == '同条件养护':
        temp = parse_temperature(ss.get('lei_ji_yang_hu_wen_du'))
        if temp is None: probs.append("累计温度解析失败:%r" % ss.get('lei_ji_yang_hu_wen_du'))
    if not t['samples']: probs.append("台账无样品编号")
    if t['temp'] is None: probs.append("台账无温度")
    if t['hum'] is None: probs.append("台账无湿度")
    rec = (code, len(t['samples']), gui, grade, vol, yh, temp, t['temp'], t['hum'], probs)
    (ok if not probs else warn).append(rec if probs else rec)
    if probs:
        print("%3d %s ✗ %s" % (i, code[-5:], "; ".join(probs)))

print()
print("=" * 60)
print("可自动跑: %d 份 | 有问题: %d 份 | 合计 %d" % (len(ok), len(warn), len(tasks)))
print("养护分布:", dict(Counter(r[5] for r in ok + warn)))
print("规格分布:", dict(Counter(r[2] for r in ok + warn)))
print("等级分布:", dict(Counter(r[3] for r in ok + warn)))
print()
print("--- 备注A 生成样例 ---")
for r in ok[:6]:
    print("  %s  %s" % (r[0][-5:], repr(format_remark_a(r[4], r[6]))))
print("--- 同条件样例 ---")
n = 0
for r in ok:
    if r[5] == '同条件养护' and n < 4:
        n += 1; print("  %s  %s" % (r[0][-5:], repr(format_remark_a(r[4], r[6]))))
