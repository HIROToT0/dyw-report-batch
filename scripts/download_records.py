# -*- coding: utf-8 -*-
"""批量下载【原始记录】或【电子报告】文件。

接口(逆向自编制页 downWordPdf):
    GET /ptcompile/downloadOfficialReport.do
        ?reportCode=<完整编号>&reportId=551&fileType=pdf|word&reportType=record|report

两个必须记住的参数：
  - reportId 一定要传【模板 ID 551】(TEMPLATE_ID)，传记录级 ID 会报"报告编制数据未保存！"
  - reportType 是字符串: record=原始记录 / report=电子报告

用法:
    python download_records.py --codes DWYHY20260000009096,DWYHY20260000009965
    python download_records.py                      # 按台账(DYW_LEDGER)全部编号下载
    python download_records.py --kind report        # 要电子报告
    python download_records.py --type word          # 要 word 版
"""
import argparse, csv, os, sys, time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dyw_client import DywClient, BASE

TEMPLATE_ID = "551"          # 模板 ID，绝不能换成记录级 reportId
KIND = {"record": "原始记录", "report": "电子报告"}
EXT = {"pdf": ".pdf", "word": ".docx"}


def codes_from_ledger(path):
    """从台账 xlsx 的『任务列表』sheet 读报告编号（A 列），去重保序。"""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[
        "任务列表" if "任务列表" in openpyxl.load_workbook(path, data_only=True).sheetnames else 0]
    out, seen = [], set()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if not v:
            continue
        c = str(v).strip()
        if c and c not in seen:
            seen.add(c)
            out.append(c)
    return out


def download(c, code, kind, ftype, dest):
    url = ("%s/ptcompile/downloadOfficialReport.do"
           "?reportCode=%s&reportId=%s&fileType=%s&reportType=%s"
           % (BASE, code, TEMPLATE_ID, ftype, kind))
    r = c.session.get(url, timeout=90)
    b = r.content
    if r.status_code != 200:
        return False, "HTTP %s" % r.status_code, 0
    head = b[:4]
    if not (head == b"%PDF" or head[:2] == b"PK"):
        import re
        m = re.search(rb"<u>(.*?)</u>", b)
        msg = m.group(1).decode("utf-8", "ignore") if m else b[:80].decode("utf-8", "ignore")
        return False, msg.strip(), len(b)
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    with open(dest, "wb") as f:
        f.write(b)
    return True, "", len(b)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--codes", default="", help="逗号分隔的报告编号；不给则读台账")
    ap.add_argument("--ledger", default=os.getenv("DYW_LEDGER", ""), help="台账 xlsx 路径")
    ap.add_argument("--kind", default="record", choices=["record", "report"], help="record=原始记录(默认)")
    ap.add_argument("--type", default="pdf", choices=["pdf", "word"])
    ap.add_argument("--out", default=os.path.join(os.getcwd(), "outputs"))
    ap.add_argument("--retry", type=int, default=2)
    a = ap.parse_args()

    if a.codes:
        codes = [x.strip() for x in a.codes.replace("\n", ",").split(",") if x.strip()]
    elif a.ledger and os.path.exists(a.ledger):
        codes = codes_from_ledger(a.ledger)
    else:
        print("请给 --codes，或设置 DYW_LEDGER 指向台账 xlsx")
        return 1
    print("待下载: %d 份 (%s / %s)" % (len(codes), KIND[a.kind], a.type))

    c = DywClient(); c.login()
    folder = os.path.join(a.out, "%s_%s" % (KIND[a.kind], a.type))
    rows, fails = [], []
    t0 = time.time()
    for i, code in enumerate(codes, 1):
        dest = os.path.join(folder, "%s_%s%s" % (code, KIND[a.kind], EXT[a.type]))
        ok, msg, size = False, "", 0
        for _ in range(a.retry + 1):
            try:
                ok, msg, size = download(c, code, a.kind, a.type, dest)
            except Exception as e:
                ok, msg = False, str(e)[:60]
            if ok:
                break
            time.sleep(1.5)
        rows.append([code, os.path.basename(dest), size, "" if ok else msg])
        if ok:
            print("%3d/%d %s ✓ %d KB" % (i, len(codes), code[-5:], size // 1024))
        else:
            fails.append((code, msg))
            print("%3d/%d %s ✗ %s" % (i, len(codes), code[-5:], msg))

    csv_path = os.path.join(a.out, "下载清单_%s_%s.csv" % (KIND[a.kind], a.type))
    os.makedirs(a.out, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["报告编号", "文件名", "字节数", "备注"]); w.writerows(rows)

    print("\n" + "=" * 58)
    print("成功 %d / %d，失败 %d，耗时 %.1f 秒" % (len(codes) - len(fails), len(codes),
                                                len(fails), time.time() - t0))
    print("输出目录:", folder)
    if fails:
        for code, msg in fails:
            print("   失败 %s %s" % (code, msg))
    return 0


if __name__ == "__main__":
    sys.exit(main())
